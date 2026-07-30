from __future__ import annotations

import ast
import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
TOOL_PATH = ROOT / "tools" / "arm2_assemble_ftas_quarantine.py"
JTA_TOOL_PATH = ROOT / "tools" / "arm2_assemble_jta_quarantine.py"


def _load_module(path: Path, name: str):
    scripts = str(ROOT / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    spec = importlib.util.spec_from_file_location(
        name, path
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _load_tool():
    return _load_module(TOOL_PATH, "arm2_assemble_ftas_quarantine_test")


def _load_jta_tool():
    return _load_module(JTA_TOOL_PATH, "arm2_assemble_jta_quarantine_test")


def _write_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def _write_seen_panel(path: Path) -> None:
    rows = []
    for year in range(2018, 2026):
        vintage = "preliminary" if year == 2025 else "confirmed"
        for month in range(1, 13):
            rows.append({
                "pref_code": "18",
                "pref_name": "福井県",
                "year": year,
                "month": month,
                "total_stays": 1.0,
                "foreign_stays_10plus": 1.0,
                "vintage": vintage,
            })
    pd.DataFrame(rows).to_csv(path, index=False)


def _write_jta_workbook(path: Path, months: list[int]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if not months:
        wb.create_sheet("Sheet")
    for month in months:
        for title in (f"第2表({month}月)", f"参考第1表({month}月)"):
            ws = wb.create_sheet(title)
            ws["A5"] = "　18福井県"
            ws["B5"] = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _build_s2_ready_panel(include_2026: bool) -> pd.DataFrame:
    rows = []
    for year in range(2018, 2025):
        for month in range(1, 13):
            rows.append({
                "pref_code": "18",
                "pref_name": "福井県",
                "year": year,
                "month": month,
                "total_stays": 1.0,
                "foreign_stays_10plus": 1.0,
                "vintage": "confirmed",
            })
    for month in range(1, 13):
        rows.append({
            "pref_code": "18",
            "pref_name": "福井県",
            "year": 2025,
            "month": month,
            "total_stays": 1.0,
            "foreign_stays_10plus": 1.0,
            "vintage": "confirmed",
        })
    if include_2026:
        rows.append({
            "pref_code": "18",
            "pref_name": "福井県",
            "year": 2026,
            "month": 1,
            "total_stays": 1.0,
            "foreign_stays_10plus": 1.0,
            "vintage": "preliminary",
        })
    return pd.DataFrame(rows)


def test_ftas_assembly_source_guardrails() -> None:
    source = TOOL_PATH.read_text(encoding="utf-8")
    assert "data/quarantine/arm2" not in source
    assert "QUARANTINE_ROOT" not in source
    assert "urllib" not in source
    assert "requests" not in source
    assert "httpx" not in source
    assert "import pandas" not in source
    assert "from pandas" not in source
    assert "pd.read_csv" not in source
    assert "read_csv(" not in source
    assert '"clone"' not in source
    assert '"fetch"' not in source
    assert '"pull"' not in source


def test_tool_imports_without_pandas() -> None:
    tool = _load_tool()
    assert "pandas" not in tool.__dict__


def test_seen_merged_digest_mismatch_raises(tmp_path, monkeypatch) -> None:
    tool = _load_tool()
    pinned = tmp_path / "pinned"
    target = tmp_path / "ftas"
    _write_bytes(pinned / "merged_survey_2023.csv", b"seen-2023")

    monkeypatch.setattr(tool.arm2_quarantine, "PINNED_MERGED_DIR", pinned)
    monkeypatch.setattr(tool.arm2_quarantine, "FTAS_DIR", target)
    monkeypatch.setitem(
        tool.arm2_quarantine.EXPECTED_SEEN_MERGED_SHA256,
        2023,
        "0" * 64,
    )

    with pytest.raises(tool.ToolError, match="checksum mismatch"):
        tool._copy_pinned_merged_waves()


def test_prefix_mismatch_raises(tmp_path) -> None:
    tool = _load_tool()
    reference = tmp_path / "merged_survey_2026.csv"
    extended = tmp_path / "extended.csv"
    _write_bytes(reference, b"abc123")
    _write_bytes(extended, b"abc124suffix")

    with pytest.raises(tool.ToolError, match="human decision"):
        tool._verify_2026_prefix(reference, extended)


def test_non_empty_target_directory_raises(tmp_path, monkeypatch) -> None:
    tool = _load_tool()
    target = tmp_path / "ftas"
    target.mkdir()
    _write_bytes(target / "existing.txt", b"x")
    monkeypatch.setattr(tool.arm2_quarantine, "FTAS_DIR", target)

    with pytest.raises(tool.ToolError, match="must be empty"):
        tool._ensure_empty_target()


def test_jta_assembly_source_guardrails() -> None:
    source = JTA_TOOL_PATH.read_text(encoding="utf-8")
    assert "data/quarantine/arm2" not in source
    assert "QUARANTINE_ROOT" not in source
    for token in ("urllib", "requests", "httpx", "socket", "urlopen"):
        assert token not in source

    tree = ast.parse(source)
    print_calls = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call):
            continue
        if not isinstance(node.func, ast.Name) or node.func.id != "print":
            continue
        if any(keyword.arg == "file" for keyword in node.keywords):
            continue
        segment = ast.get_source_segment(source, node)
        assert segment is not None
        print_calls.append(segment)

    assert any(
        "metadata['filename']" in segment and "metadata['sha256']" in segment
        for segment in print_calls
    )
    assert any(
        "output_path" in segment and "output_sha" in segment and "row_count" in segment
        for segment in print_calls
    )
    assert any("labels" in segment for segment in print_calls)
    for forbidden in ("total_stays", "foreign_stays_10plus", "sum(", "min(", "max("):
        assert all(forbidden not in segment for segment in print_calls)


def test_empty_monthly_parse_raises(tmp_path) -> None:
    tool = _load_jta_tool()
    workbook = tmp_path / "blank_2026.xlsx"
    _write_jta_workbook(workbook, [])

    with pytest.raises(tool.ToolError, match="yielded no rows"):
        tool._parse_source_workbook(
            workbook,
            year=2026,
            vintage="preliminary",
            require_complete_year=False,
        )


def test_duplicate_monthly_files_raise(tmp_path) -> None:
    tool = _load_jta_tool()
    january_a = tmp_path / "2026-01-a.xlsx"
    january_b = tmp_path / "2026-01-b.xlsx"
    _write_jta_workbook(january_a, [1])
    _write_jta_workbook(january_b, [1])

    with pytest.raises(tool.ToolError, match="duplicate 2026 months"):
        tool._parse_monthly_workbooks([january_a, january_b])


def test_missing_confirmed_year_raises(tmp_path) -> None:
    tool = _load_jta_tool()
    workbook = tmp_path / "2025-partial.xlsx"
    _write_jta_workbook(workbook, [1])

    with pytest.raises(tool.ToolError, match="all 12 months"):
        tool._parse_source_workbook(
            workbook,
            year=2025,
            vintage="confirmed",
            require_complete_year=True,
        )


def test_no_2026_row_raises_from_s2_gate() -> None:
    tool = _load_jta_tool()
    panel = _build_s2_ready_panel(include_2026=False)

    with pytest.raises(tool.ToolError, match="2026"):
        tool._validate_s2_contract(panel)


def test_non_empty_jta_target_directory_raises(tmp_path, monkeypatch) -> None:
    tool = _load_jta_tool()
    target = tmp_path / "jta"
    target.mkdir()
    _write_bytes(target / "existing.txt", b"x")
    monkeypatch.setattr(tool.arm2_quarantine, "JTA_DIR", target)

    with pytest.raises(tool.ToolError, match="must be empty"):
        tool._ensure_empty_target()


def test_s2_oracle_result_is_discarded(monkeypatch) -> None:
    tool = _load_jta_tool()
    panel = _build_s2_ready_panel(include_2026=True)

    class Sentinel:
        def __bool__(self):
            raise AssertionError("result should be discarded")

        def __getitem__(self, key):
            raise AssertionError(f"result should not be indexed: {key}")

        def __iter__(self):
            raise AssertionError("result should not be iterated")

        def __repr__(self):
            raise AssertionError("result should not be rendered")

    monkeypatch.setattr(tool, "build_s2_descriptive_report", lambda frame: Sentinel())

    assert tool._validate_s2_contract(panel) is None


def test_jta_main_writes_manifest_and_frame_basis(tmp_path, monkeypatch) -> None:
    tool = _load_jta_tool()
    seen_panel = tmp_path / "accommodation_panel.csv"
    confirmed = tmp_path / "2025_confirmed.xlsx"
    monthly = tmp_path / "2026_01.xlsx"
    target = tmp_path / "jta"
    _write_seen_panel(seen_panel)
    _write_jta_workbook(confirmed, list(range(1, 13)))
    _write_jta_workbook(monthly, [1])

    monkeypatch.setattr(tool, "SEEN_PANEL_PATH", seen_panel)
    monkeypatch.setattr(tool.arm2_quarantine, "JTA_DIR", target)

    assert tool.main(
        [
            "--confirmed-2025",
            str(confirmed),
            "--monthly-2026",
            str(monthly),
        ]
    ) == 0

    written = pd.read_csv(target / tool.DESTINATION_FILENAME, dtype={"pref_code": str})
    assert list(written.columns) == list(tool.BASE_COLUMNS) + [tool.FRAME_BASIS_COLUMN]
    assert not written.duplicated(["pref_code", "year", "month"]).any()
    rows_2025 = written[written["year"] == 2025]
    assert len(rows_2025) == 12
    assert rows_2025["vintage"].eq("confirmed").all()
    manifest = json.loads((target / tool.MANIFEST_FILENAME).read_text(encoding="utf-8"))
    assert manifest["frame_basis_included"] is True
